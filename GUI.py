import tkinter as tk
from tkinter import filedialog, messagebox
import openpyxl
from copy import copy

def main_func(filepath):
    print(f"Processing file: {filepath}")
    wb = openpyxl.load_workbook(filepath)
    ws = wb['sheet_a']
    header = [cell.value for cell in next(ws.iter_rows(max_row=1))]
    
    key_col_idx = header.index('oznaka') + 1
    keys = set(
        cell.value 
        for col in ws.iter_cols(min_row=2, min_col=key_col_idx, max_col=key_col_idx)
        for cell in col if cell.value is not None
    )

    for key in keys:
        export_workbook = openpyxl.Workbook()
        export_sheet = export_workbook.active
        
        for col_idx, cell in enumerate(next(ws.iter_rows(max_row=1)), start=1):
            header_cell = export_sheet.cell(row=1, column=col_idx, value=cell.value)
            if cell.has_style:
                header_cell.font = copy(cell.font)
                header_cell.border = copy(cell.border)
                header_cell.fill = copy(cell.fill)
                header_cell.number_format = copy(cell.number_format)
                header_cell.protection = copy(cell.protection)
                header_cell.alignment = copy(cell.alignment)

        row_idx = 2
        for row in ws.iter_rows(min_row=2):
            if row[key_col_idx - 1].value == key:
                for col_idx, cell in enumerate(row, start=1):
                    dest_cell = export_sheet.cell(row=row_idx, column=col_idx, value=cell.value)
                    if cell.has_style:
                        dest_cell.font = copy(cell.font)
                        dest_cell.border = copy(cell.border)
                        dest_cell.fill = copy(cell.fill)
                        dest_cell.number_format = copy(cell.number_format)
                        dest_cell.protection = copy(cell.protection)
                        dest_cell.alignment = copy(cell.alignment)
                row_idx += 1
        
        export_filename = f"C:/Users/sinis/Documents/{key}_export_test1.xlsx"
        export_workbook.save(export_filename)
        print(f"Exported {key} to {export_filename}")



def open_file_dialog():
    filepath = filedialog.askopenfilename(
        title="Select an Excel File",
        filetypes=(("Excel files", "*.xlsx"), ("All files", "*.*"))
    )
    if filepath:
        main_func(filepath)
        messagebox.showinfo("Success", "Workbook processed successfully!")
        root.destroy()

# Create the main GUI window
root = tk.Tk()
root.title("Excel Workbook Processor")

# Create the buttons
frame = tk.Frame(root)
frame.pack(pady=20)

load_button = tk.Button(frame, text="Load Workbook", command=open_file_dialog, width=20)
load_button.pack(side=tk.LEFT, padx=10)

exit_button = tk.Button(frame, text="Exit", command=root.destroy, width=20)
exit_button.pack(side=tk.LEFT, padx=10)

root.mainloop()
